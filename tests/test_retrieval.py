from __future__ import annotations

import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from docx import Document
from openpyxl import Workbook

from agent_workflow_studio.retrieval.chunking import chunk_text
from agent_workflow_studio.retrieval.extractors import FilePayload, decode_text, extract_text_from_file
from agent_workflow_studio.retrieval.rag import TextSource, build_index, retrieve


class FakeEmbedder:
    def embed(self, texts):
        return [[float(text.lower().count("robot")), float(text.lower().count("battery")), 1.0] for text in texts]


class RetrievalTests(unittest.TestCase):
    def test_decode_cp949(self) -> None:
        value = "test korean text".encode("cp949")
        self.assertEqual(decode_text(value), "test korean text")

    def test_text_extraction_and_chunking(self) -> None:
        payload = FilePayload("notes.txt", b"alpha beta gamma delta epsilon")
        self.assertEqual(extract_text_from_file(payload), "alpha beta gamma delta epsilon")
        chunks = chunk_text("one two three four five six", chunk_size=10, overlap=2)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunks))

    def test_docx_extraction(self) -> None:
        doc = Document()
        doc.add_paragraph("hello")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "a"
        table.cell(0, 1).text = "b"
        buffer = io.BytesIO()
        doc.save(buffer)
        text = extract_text_from_file(FilePayload("sample.docx", buffer.getvalue()))
        self.assertIn("hello", text)
        self.assertIn("a | b", text)

    def test_xlsx_extraction(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["name", "value"])
        ws.append(["robot", 10])
        buffer = io.BytesIO()
        wb.save(buffer)
        text = extract_text_from_file(FilePayload("sample.xlsx", buffer.getvalue()))
        self.assertIn("[Sheet: Data]", text)
        self.assertIn("robot | 10", text)

    def test_rag_is_ui_independent(self) -> None:
        embedder = FakeEmbedder()
        index = build_index([TextSource("s1", "robot", "file", "robot robot control"), TextSource("s2", "battery", "notion", "battery battery cell")], embedder)
        hits = retrieve("robot", index, embedder, top_k=1)
        self.assertEqual(hits[0].source_id, "s1")


if __name__ == "__main__":
    unittest.main()
