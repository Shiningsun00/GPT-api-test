from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from agent_workflow_studio.core.errors import FileExtractionError, UnsupportedFileTypeError

SUPPORTED_SUFFIXES = frozenset({".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".txt", ".md"})


@dataclass(frozen=True, slots=True)
class FilePayload:
    name: str
    data: bytes

    @property
    def sha256(self) -> str:
        return sha256_bytes(self.data)

    @property
    def size(self) -> int:
        return len(self.data)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def extract_text_from_file(file_payload: FilePayload) -> str:
    name = Path(file_payload.name).name
    suffix = Path(name).suffix.lower()
    data = file_payload.data
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedFileTypeError(f"unsupported file type: {suffix or '<none>'}")
    try:
        if suffix in {".txt", ".md", ".csv"}:
            return decode_text(data)
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            return "\n\n".join((page.extract_text() or "") for page in reader.pages)
        if suffix == ".docx":
            doc = Document(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        paragraphs.append(" | ".join(cells))
            return "\n".join(paragraphs)
        if suffix == ".xlsx":
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines: list[str] = []
            for worksheet in workbook.worksheets:
                lines.append(f"[Sheet: {worksheet.title}]")
                for row in worksheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(value.strip() for value in values):
                        lines.append(" | ".join(values))
            return "\n".join(lines)
        if suffix == ".pptx":
            presentation = Presentation(io.BytesIO(data))
            lines: list[str] = []
            for index, slide in enumerate(presentation.slides, start=1):
                lines.append(f"[Slide {index}]")
                for shape in slide.shapes:
                    text = getattr(shape, "text", "")
                    if isinstance(text, str) and text.strip():
                        lines.append(text.strip())
            return "\n".join(lines)
    except (UnsupportedFileTypeError, FileExtractionError):
        raise
    except Exception as exc:
        raise FileExtractionError(f"failed to extract text from {name}: {exc}") from exc
    raise UnsupportedFileTypeError(f"unsupported file type: {suffix}")
